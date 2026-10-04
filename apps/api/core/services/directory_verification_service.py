import csv
import io
import json
import re
from dataclasses import dataclass
from decimal import Decimal
from typing import Iterable, Optional

from django.db import transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook

from core.models import DirectoryUser, DirectoryVerifyJob, DirectoryVerifyJobStatus, DirectoryVerifyRow, DirectoryVerifyVerdict


class DirectoryVerificationError(Exception):
    def __init__(self, message: str, code: str = "directory_verification_error"):
        super().__init__(message)
        self.code = code


def normalize_email(value: str) -> str:
    return (value or "").strip().lower()


def normalize_name(value: str) -> str:
    cleaned = re.sub(r"\s+", " ", (value or "").strip())
    return cleaned.lower()


def normalize_phone(value: str) -> str:
    digits = re.sub(r"\D+", "", (value or ""))
    if not digits:
        return ""
    if len(digits) > 10:
        digits = digits[-10:]
    return digits


def normalize_roll(value: str) -> str:
    return re.sub(r"\s+", "", (value or "").strip().lower())


def _header_key(value: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", (value or "").strip().lower()).strip()


CANONICAL_FIELDS = [
    "name",
    "roll_no",
    "email_address",
    "phone_number",
    "official_email_issued",
    "email_pmc",
]


HEADER_ALIASES = {
    "name": {"name", "student name", "full name"},
    "roll_no": {"roll no", "roll number", "rollno", "roll"},
    "email_address": {"email", "email address", "email id", "personal email", "primary email"},
    "phone_number": {"phone", "phone number", "mobile", "mobile number", "contact", "contact number", "phone no"},
    "official_email_issued": {
        "official email issued",
        "official email",
        "institutional email",
        "institution email",
        "official email id",
    },
    "email_pmc": {"email pmc", "pmc email", "email_pmc"},
}


def suggest_mapping(headers: list[str]) -> dict:
    keyed_headers = {_header_key(h): h for h in headers}
    mapping = {}
    for canonical in CANONICAL_FIELDS:
        selected = ""
        for alias in HEADER_ALIASES.get(canonical, set()):
            if alias in keyed_headers:
                selected = keyed_headers[alias]
                break
        mapping[canonical] = selected
    return mapping


def parse_mapping(raw_mapping: object) -> dict:
    if isinstance(raw_mapping, str):
        try:
            raw_mapping = json.loads(raw_mapping)
        except json.JSONDecodeError as exc:
            raise DirectoryVerificationError("Invalid JSON for mapping.", "invalid_mapping_json") from exc
    if not isinstance(raw_mapping, dict):
        raise DirectoryVerificationError("mapping must be an object.", "invalid_mapping")

    mapping = {}
    for key in CANONICAL_FIELDS:
        value = raw_mapping.get(key, "")
        mapping[key] = str(value).strip() if value is not None else ""
    return mapping


def parse_tabular_file(file_name: str, file_bytes: bytes, sheet_name: str = "") -> tuple[list[str], list[dict], str, list[str]]:
    lower_name = (file_name or "").lower()
    if lower_name.endswith((".csv", ".tsv", ".txt")):
        try:
            text = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DirectoryVerificationError("Could not decode file as UTF-8.", "invalid_file_encoding") from exc
        lines = [line for line in text.splitlines() if line.strip()]
        if not lines:
            raise DirectoryVerificationError("Uploaded file is empty.", "empty_file")

        sample = "\n".join(lines[:5])
        delimiter = ","
        if lower_name.endswith(".tsv"):
            delimiter = "\t"
        else:
            try:
                dialect = csv.Sniffer().sniff(sample, delimiters=",\t;")
                delimiter = dialect.delimiter
            except csv.Error:
                delimiter = "\t" if sample.count("\t") > sample.count(",") else ","

        reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
        if not reader.fieldnames:
            raise DirectoryVerificationError("Uploaded file has no headers.", "missing_headers")
        headers = [str(h or "").strip() for h in reader.fieldnames]
        rows = [{h: str(v or "").strip() for h, v in row.items()} for row in reader]
        return headers, rows, "", []

    if lower_name.endswith((".xlsx", ".xlsm")):
        try:
            workbook = load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        except Exception as exc:
            raise DirectoryVerificationError("Could not read Excel workbook.", "invalid_excel_file") from exc

        worksheet_names = workbook.sheetnames
        if not worksheet_names:
            raise DirectoryVerificationError("Workbook has no worksheets.", "empty_workbook")

        worksheet = None
        selected_sheet = ""
        if sheet_name:
            if sheet_name not in worksheet_names:
                raise DirectoryVerificationError("Selected sheet was not found in workbook.", "invalid_sheet")
            worksheet = workbook[sheet_name]
            selected_sheet = sheet_name
        else:
            for candidate in workbook.worksheets:
                values = list(candidate.iter_rows(values_only=True))
                if any(any(cell not in (None, "") for cell in row) for row in values):
                    worksheet = candidate
                    selected_sheet = candidate.title
                    break
        if worksheet is None:
            raise DirectoryVerificationError("Workbook contains no data rows.", "empty_workbook")

        rows_raw = list(worksheet.iter_rows(values_only=True))
        if not rows_raw:
            raise DirectoryVerificationError("Selected worksheet has no rows.", "empty_sheet")

        headers = [str(cell or "").strip() for cell in rows_raw[0]]
        if not any(headers):
            raise DirectoryVerificationError("Selected worksheet has empty headers.", "missing_headers")

        rows = []
        for raw in rows_raw[1:]:
            row_dict = {}
            for idx, header in enumerate(headers):
                if not header:
                    continue
                value = raw[idx] if idx < len(raw) else ""
                row_dict[header] = "" if value is None else str(value).strip()
            if any(row_dict.values()):
                rows.append(row_dict)

        return headers, rows, selected_sheet, worksheet_names

    raise DirectoryVerificationError(
        "Unsupported file type. Upload CSV, TSV, XLSX, or XLSM.",
        "unsupported_file_type",
    )


def extract_row_input(row: dict, mapping: dict) -> dict:
    extracted = {}
    for canonical, source_header in mapping.items():
        value = row.get(source_header, "") if source_header else ""
        extracted[canonical] = str(value or "").strip()
    return extracted


@dataclass
class MatchOutcome:
    verdict: str
    basis: str
    confidence: Optional[Decimal]
    notes: str
    matched_user: Optional[DirectoryUser] = None
    matched_email: str = ""
    matched_name: str = ""


class DirectoryVerificationService:
    def _build_indexes(self):
        users = list(
            DirectoryUser.objects.all().only(
                "id",
                "primary_email",
                "full_name",
                "normalized_email",
                "normalized_full_name",
                "normalized_phone",
                "aliases_json",
                "phones_json",
                "roll_number",
                "external_identifier",
                "employee_id",
            )
        )

        email_index: dict[str, list[DirectoryUser]] = {}
        alias_index: dict[str, list[DirectoryUser]] = {}
        name_phone_index: dict[tuple[str, str], list[DirectoryUser]] = {}
        name_roll_index: dict[tuple[str, str], list[DirectoryUser]] = {}
        name_index: dict[str, list[DirectoryUser]] = {}

        for user in users:
            email = normalize_email(user.normalized_email or user.primary_email)
            if email:
                email_index.setdefault(email, []).append(user)

            aliases = user.aliases_json if isinstance(user.aliases_json, list) else []
            for alias in aliases:
                normalized_alias = normalize_email(str(alias))
                if normalized_alias:
                    alias_index.setdefault(normalized_alias, []).append(user)

            full_name = normalize_name(user.normalized_full_name or user.full_name)
            if full_name:
                name_index.setdefault(full_name, []).append(user)

            phones = user.phones_json if isinstance(user.phones_json, list) else []
            normalized_phones = [normalize_phone(str(p)) for p in phones]
            if user.normalized_phone:
                normalized_phones.append(normalize_phone(user.normalized_phone))
            normalized_phones = [p for p in normalized_phones if p]
            for phone in set(normalized_phones):
                if full_name:
                    name_phone_index.setdefault((full_name, phone), []).append(user)

            roll_candidates = [
                normalize_roll(user.roll_number),
                normalize_roll(user.external_identifier),
                normalize_roll(user.employee_id),
            ]
            for roll in {r for r in roll_candidates if r}:
                if full_name:
                    name_roll_index.setdefault((full_name, roll), []).append(user)

        return {
            "email_index": email_index,
            "alias_index": alias_index,
            "name_phone_index": name_phone_index,
            "name_roll_index": name_roll_index,
            "name_index": name_index,
        }

    def _is_row_usable(self, extracted: dict) -> bool:
        official_email = normalize_email(extracted.get("official_email_issued", ""))
        other_emails = [
            normalize_email(extracted.get("email_address", "")),
            normalize_email(extracted.get("email_pmc", "")),
        ]
        any_email = official_email or any(other_emails)
        name = normalize_name(extracted.get("name", ""))
        phone = normalize_phone(extracted.get("phone_number", ""))
        roll_no = normalize_roll(extracted.get("roll_no", ""))

        return bool(
            any_email
            or (name and phone)
            or (name and roll_no)
            or name
        )

    def _unique_users(self, users: Iterable[DirectoryUser]) -> list[DirectoryUser]:
        by_id = {}
        for user in users:
            by_id[user.id] = user
        return list(by_id.values())

    def _match_row(self, extracted: dict, indexes: dict) -> MatchOutcome:
        if not self._is_row_usable(extracted):
            return MatchOutcome(
                verdict=DirectoryVerifyVerdict.INVALID_INPUT,
                basis="invalid_input",
                confidence=None,
                notes="Row is missing required searchable fields.",
            )

        official_email = normalize_email(extracted.get("official_email_issued", ""))
        email_address = normalize_email(extracted.get("email_address", ""))
        email_pmc = normalize_email(extracted.get("email_pmc", ""))
        email_candidates = [e for e in [official_email, email_address, email_pmc] if e]
        name = normalize_name(extracted.get("name", ""))
        phone = normalize_phone(extracted.get("phone_number", ""))
        roll_no = normalize_roll(extracted.get("roll_no", ""))

        for email in email_candidates:
            matched = indexes["email_index"].get(email, [])
            unique = self._unique_users(matched)
            if len(unique) == 1:
                basis = "exact_official_email" if email == official_email and official_email else "exact_primary_email"
                user = unique[0]
                return MatchOutcome(
                    verdict=DirectoryVerifyVerdict.EXISTS,
                    basis=basis,
                    confidence=Decimal("1.00"),
                    notes="Matched by exact primary email.",
                    matched_user=user,
                    matched_email=user.primary_email,
                    matched_name=user.full_name,
                )
            if len(unique) > 1:
                return MatchOutcome(
                    verdict=DirectoryVerifyVerdict.AMBIGUOUS,
                    basis="exact_primary_email",
                    confidence=Decimal("0.40"),
                    notes=f"Multiple users matched email {email}.",
                )

        alias_matches = []
        for email in email_candidates:
            alias_matches.extend(indexes["alias_index"].get(email, []))
        alias_unique = self._unique_users(alias_matches)
        if len(alias_unique) == 1:
            user = alias_unique[0]
            return MatchOutcome(
                verdict=DirectoryVerifyVerdict.EXISTS,
                basis="alias_email",
                confidence=Decimal("0.95"),
                notes="Matched by alias email.",
                matched_user=user,
                matched_email=user.primary_email,
                matched_name=user.full_name,
            )
        if len(alias_unique) > 1:
            return MatchOutcome(
                verdict=DirectoryVerifyVerdict.AMBIGUOUS,
                basis="alias_email",
                confidence=Decimal("0.40"),
                notes="Multiple users matched alias email.",
            )

        if name and phone:
            candidates = self._unique_users(indexes["name_phone_index"].get((name, phone), []))
            if len(candidates) == 1:
                user = candidates[0]
                return MatchOutcome(
                    verdict=DirectoryVerifyVerdict.EXISTS,
                    basis="name_phone",
                    confidence=Decimal("0.80"),
                    notes="Matched by full name + phone.",
                    matched_user=user,
                    matched_email=user.primary_email,
                    matched_name=user.full_name,
                )
            if len(candidates) > 1:
                return MatchOutcome(
                    verdict=DirectoryVerifyVerdict.AMBIGUOUS,
                    basis="name_phone",
                    confidence=Decimal("0.35"),
                    notes="Multiple users matched full name + phone.",
                )

        if name and roll_no:
            candidates = self._unique_users(indexes["name_roll_index"].get((name, roll_no), []))
            if len(candidates) == 1:
                user = candidates[0]
                return MatchOutcome(
                    verdict=DirectoryVerifyVerdict.EXISTS,
                    basis="name_roll_no",
                    confidence=Decimal("0.75"),
                    notes="Matched by full name + roll/identifier.",
                    matched_user=user,
                    matched_email=user.primary_email,
                    matched_name=user.full_name,
                )
            if len(candidates) > 1:
                return MatchOutcome(
                    verdict=DirectoryVerifyVerdict.AMBIGUOUS,
                    basis="name_roll_no",
                    confidence=Decimal("0.35"),
                    notes="Multiple users matched full name + roll/identifier.",
                )

        if name:
            candidates = self._unique_users(indexes["name_index"].get(name, []))
            if len(candidates) == 1:
                user = candidates[0]
                return MatchOutcome(
                    verdict=DirectoryVerifyVerdict.EXISTS,
                    basis="name_only",
                    confidence=Decimal("0.55"),
                    notes="Low-confidence full-name-only match.",
                    matched_user=user,
                    matched_email=user.primary_email,
                    matched_name=user.full_name,
                )
            if len(candidates) > 1:
                return MatchOutcome(
                    verdict=DirectoryVerifyVerdict.AMBIGUOUS,
                    basis="name_only",
                    confidence=Decimal("0.20"),
                    notes="Multiple users matched full name only.",
                )

        return MatchOutcome(
            verdict=DirectoryVerifyVerdict.DOES_NOT_EXIST,
            basis="no_match",
            confidence=Decimal("0.00"),
            notes="No matching directory user found.",
        )

    @transaction.atomic
    def run_job(
        self,
        *,
        created_by,
        source_filename: str,
        source_type: str,
        sheet_name: str,
        rows: list[dict],
        mapping: dict,
    ) -> DirectoryVerifyJob:
        job = DirectoryVerifyJob.objects.create(
            created_by=created_by,
            source_filename=source_filename,
            source_type=source_type,
            status=DirectoryVerifyJobStatus.PROCESSING,
            row_count=len(rows),
            column_mapping_json=mapping,
            sheet_name=sheet_name or "",
        )

        indexes = self._build_indexes()
        rows_to_create = []
        summary = {
            "total_rows": len(rows),
            "exists": 0,
            "does_not_exist": 0,
            "ambiguous": 0,
            "invalid_input": 0,
        }

        for idx, row in enumerate(rows, start=2):
            extracted = extract_row_input(row, mapping)
            outcome = self._match_row(extracted, indexes)
            summary[outcome.verdict] += 1
            rows_to_create.append(
                DirectoryVerifyRow(
                    job=job,
                    row_no=idx,
                    input_json=extracted,
                    matched_directory_user=outcome.matched_user,
                    matched_email=outcome.matched_email,
                    matched_name=outcome.matched_name,
                    match_basis=outcome.basis,
                    verdict=outcome.verdict,
                    confidence_score=outcome.confidence,
                    notes=outcome.notes,
                )
            )

        DirectoryVerifyRow.objects.bulk_create(rows_to_create, batch_size=500)
        job.status = DirectoryVerifyJobStatus.COMPLETED
        job.summary_json = summary
        job.completed_at = timezone.now()
        job.save(update_fields=["status", "summary_json", "completed_at", "updated_at"])
        return job


EXPORT_COLUMNS = [
    "row_no",
    "name",
    "roll_no",
    "phone_number",
    "email_address",
    "official_email_issued",
    "email_pmc",
    "matched_email",
    "matched_name",
    "match_basis",
    "verdict",
    "confidence_score",
    "notes",
]


def build_export_records(rows: Iterable[DirectoryVerifyRow]) -> list[dict]:
    records = []
    for row in rows:
        input_json = row.input_json or {}
        records.append(
            {
                "row_no": row.row_no,
                "name": input_json.get("name", ""),
                "roll_no": input_json.get("roll_no", ""),
                "phone_number": input_json.get("phone_number", ""),
                "email_address": input_json.get("email_address", ""),
                "official_email_issued": input_json.get("official_email_issued", ""),
                "email_pmc": input_json.get("email_pmc", ""),
                "matched_email": row.matched_email,
                "matched_name": row.matched_name,
                "match_basis": row.match_basis,
                "verdict": row.verdict,
                "confidence_score": str(row.confidence_score) if row.confidence_score is not None else "",
                "notes": row.notes,
            }
        )
    return records


def export_records_csv(records: list[dict]) -> bytes:
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=EXPORT_COLUMNS)
    writer.writeheader()
    for row in records:
        writer.writerow(row)
    return output.getvalue().encode("utf-8")


def export_records_xlsx(records: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "verification"
    ws.append(EXPORT_COLUMNS)
    for row in records:
        ws.append([row.get(column, "") for column in EXPORT_COLUMNS])
    stream = io.BytesIO()
    wb.save(stream)
    return stream.getvalue()
